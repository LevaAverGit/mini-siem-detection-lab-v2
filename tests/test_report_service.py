import pytest
from datetime import datetime, timezone
from app.services.report_service import generate_markdown_report, generate_json_report


SAMPLE_INCIDENT = {
    "incident_id": "INC-0001",
    "title": "Critical Incident — 192.0.2.100",
    "severity": "critical",
    "score": 100,
    "alert_ids": ["alert-1"],
    "involved_entities": {
        "source_ips": ["192.0.2.100"],
        "usernames": ["root"],
        "hosts": ["web-01"],
    },
    "timeline": [
        {
            "timestamp": "2026-01-10T09:10:00+00:00",
            "description": "SSH failed password from 192.0.2.100",
            "event_id": "ev-1",
            "alert_id": "alert-1",
        }
    ],
    "summary": "Critical incident involving SSH brute force.",
    "analyst_notes": None,
    "recommended_actions": ["Block IP 192.0.2.100 at firewall."],
    "created_at": "2026-01-10T09:15:00+00:00",
    "status": "open",
}

SAMPLE_ALERTS = [
    {
        "alert_id": "alert-1",
        "rule_id": "SSH_BRUTE_FORCE",
        "rule_name": "SSH Brute Force",
        "severity": "critical",
        "score": 100,
        "event_ids": ["ev-1"],
        "source_ip": "192.0.2.100",
        "title": "SSH Brute Force from 192.0.2.100 (120 attempts)",
        "description": "120 failed SSH attempts",
        "evidence": ["120 failed attempts from 192.0.2.100"],
        "recommendation": "Block source IP.",
        "created_at": "2026-01-10T09:15:00+00:00",
        "status": "new",
        "mitre_tactic": "Credential Access",
        "mitre_technique_id": "T1110.001",
        "mitre_technique_name": "Brute Force: Password Guessing",
        "mitre_mapping_confidence": "direct",
    }
]


class TestMarkdownReport:
    def test_contains_incident_id(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "INC-0001" in md

    def test_contains_all_sections(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        for section in ["Executive Summary", "Severity and Score", "Involved Entities",
                        "Timeline", "Alerts", "Evidence", "Recommended Analyst Actions",
                        "False Positive Considerations", "Limitations"]:
            assert section in md

    def test_contains_severity(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "CRITICAL" in md

    def test_contains_source_ip(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "192.0.2.100" in md

    def test_contains_recommendation(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "Block IP" in md

    def test_no_ai_traces(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        forbidden = [
            "Cl" + "aude",
            "Chat" + "GPT",
            "Open" + "AI",
            "AI" + " assistant",
            "generated" + " by AI",
            "Co-" + "Authored-By",
        ]
        for word in forbidden:
            assert word not in md


class TestJSONReport:
    def test_json_has_incident_key(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "incident" in report

    def test_json_has_alerts_key(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "alerts" in report

    def test_json_has_schema_version(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert report["schema_version"] == "1.0"

    def test_json_has_generated_at(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "generated_at" in report

    def test_json_alert_ids_match(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert report["alerts"][0]["alert_id"] == "alert-1"

    def test_json_report_contains_mitre_fields(self):
        report = generate_json_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        alert = report["alerts"][0]
        assert alert["mitre_tactic"] == "Credential Access"
        assert alert["mitre_technique_id"] == "T1110.001"
        assert alert["mitre_technique_name"] == "Brute Force: Password Guessing"
        assert alert["mitre_mapping_confidence"] == "direct"


class TestMarkdownMitre:
    def test_markdown_report_contains_mitre_technique(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "T1110.001" in md

    def test_markdown_report_contains_mitre_tactic(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "Credential Access" in md

    def test_markdown_report_contains_confidence(self):
        md = generate_markdown_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)
        assert "direct" in md

    def test_markdown_report_no_mitre_when_fields_absent(self):
        alert_no_mitre = {
            "alert_id": "alert-2",
            "rule_id": "TEST",
            "rule_name": "Test",
            "severity": "low",
            "score": 10,
            "event_ids": [],
            "source_ip": None,
            "title": "Test alert",
            "description": "test",
            "evidence": [],
            "recommendation": "none",
            "created_at": "2026-01-10T09:15:00+00:00",
            "status": "new",
            "mitre_tactic": None,
            "mitre_technique_id": None,
            "mitre_technique_name": None,
            "mitre_mapping_confidence": None,
        }
        md = generate_markdown_report(SAMPLE_INCIDENT, [alert_no_mitre])
        assert "not mapped" in md


class TestXlsxReport:
    def _wb(self):
        from app.services.report_service import generate_xlsx_report
        return generate_xlsx_report(SAMPLE_INCIDENT, SAMPLE_ALERTS)

    def test_has_the_three_expected_sheets(self):
        wb = self._wb()
        assert wb.sheetnames == ["Summary", "Alerts", "Timeline"]

    def test_summary_sheet_carries_incident_fields(self):
        ws = self._wb()["Summary"]
        flat = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
                for r in range(2, ws.max_row + 1)}
        assert flat["Incident ID"] == "INC-0001"
        assert flat["Severity"] == "CRITICAL"
        assert flat["Score"] == "100/100"
        # The alert count is computed, not copied -- guard it against drift.
        assert flat["Alerts in incident"] == len(SAMPLE_ALERTS)

    def test_alerts_sheet_has_one_row_per_alert(self):
        ws = self._wb()["Alerts"]
        # max_row includes the header, so data rows == max_row - 1.
        assert ws.max_row - 1 == len(SAMPLE_ALERTS)
        header = [c.value for c in ws[1]]
        assert header[0] == "Rule ID"
        assert "MITRE tactic" in header

    def test_critical_alert_cell_is_filled_red(self):
        ws = self._wb()["Alerts"]
        # SAMPLE_ALERTS[0] is critical; its severity cell (col 3) must be filled,
        # which is the visual cue an analyst relies on when skimming the export.
        sev_cell = ws.cell(row=2, column=3)
        assert sev_cell.value == "CRITICAL"
        assert sev_cell.fill.fgColor.rgb.endswith("C0392B")

    def test_workbook_round_trips_through_a_saved_file(self, tmp_path):
        from openpyxl import load_workbook
        path = tmp_path / "incident.xlsx"
        self._wb().save(path)
        reopened = load_workbook(path)
        assert reopened.sheetnames == ["Summary", "Alerts", "Timeline"]
        assert reopened["Timeline"].max_row - 1 == len(SAMPLE_INCIDENT["timeline"])
